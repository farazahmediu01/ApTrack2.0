"""
Roster workbook reader.

One faculty tab holds several batch blocks stacked vertically. Blocks are found
by MARKER, never by row arithmetic - block height follows the student count, so
offsets shift the moment a batch gains a student.

Pure: reads a file, returns data, touches no network. No domain rules live here;
this layer only says what the sheet contains.
"""

from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

BLOCK_START = "attendance sheet"
HEADER_MARK = "s. no"
FOOTER_ACTUAL = "actual session covered"
FOOTER_CPC = "session covered as per cpc"
FOOTER_MODULE = "module"


@dataclass
class Student:
    row: int
    student_id: str
    name: str
    course: str
    marks: list[str]        # one entry per date column, upper-cased, "" if blank


@dataclass
class Batch:
    sheet: str
    top: int
    batch_code: str
    faculty: str
    days: str
    time: str
    month: str
    start_date: date | None
    dates: list[date] = field(default_factory=list)
    date_cols: list[int] = field(default_factory=list)
    actual: list[int | None] = field(default_factory=list)   # per date column
    planned: list[int | None] = field(default_factory=list)  # per date column
    modules: list[str] = field(default_factory=list)         # per date column
    students: list[Student] = field(default_factory=list)


def _txt(v) -> str:
    return "" if v is None else str(v).strip()


def _norm(v) -> str:
    return _txt(v).lower()


def _as_int(v):
    s = _txt(v)
    if not s:
        return None
    try:
        return int(float(s))
    except ValueError:
        return None


def _is_date(v) -> bool:
    return isinstance(v, (datetime, date))


class _Grid:
    """1-based cell access over a sheet snapshot, tolerant of ragged rows."""

    def __init__(self, ws):
        self.rows = [list(r) for r in ws.iter_rows(values_only=True)]

    def __len__(self):
        return len(self.rows)

    def at(self, r: int, c: int):
        if not (1 <= r <= len(self.rows)):
            return None
        row = self.rows[r - 1]
        return row[c - 1] if 1 <= c <= len(row) else None

    def rows_containing(self, needle: str) -> list[int]:
        return [i for i, row in enumerate(self.rows, 1)
                if any(_norm(v) == needle for v in row)]

    def col_of(self, r: int, needle: str) -> int | None:
        if not (1 <= r <= len(self.rows)):
            return None
        for c, v in enumerate(self.rows[r - 1], 1):
            if _norm(v) == needle:
                return c
        return None

    def label_row(self, lo: int, hi: int, needle: str) -> int | None:
        return next((r for r in range(lo, min(hi, len(self.rows)) + 1)
                     if self.col_of(r, needle) is not None), None)


def _paired(g: _Grid, lo: int, hi: int, label: str, raw: bool = False):
    """
    Value to the right of `label`, searching rows lo..hi.

    The gap between a label and its value is not fixed - merged cells push some
    values several columns right ('Batch Start Date' in J, its date in M) - so
    scan rightwards to the first non-empty cell rather than assuming +1.
    """
    for r in range(lo, min(hi, len(g)) + 1):
        c = g.col_of(r, label)
        if c is None:
            continue
        for step in range(1, 6):
            v = g.at(r, c + step)
            if _txt(v):
                return v if raw else _txt(v)
    return None if raw else ""


def _footer_values(g: _Grid, lo: int, hi: int, label: str, cols: list[int]) -> list:
    r = g.label_row(lo, hi, label)
    return [g.at(r, c) for c in cols] if r else [None] * len(cols)


def parse_sheet(ws) -> list[Batch]:
    g = _Grid(ws)
    starts = g.rows_containing(BLOCK_START)
    if not starts:
        return []
    edges = starts + [len(g) + 1]
    return [b for b in (_parse_block(g, ws.title, starts[i], edges[i + 1] - 1)
                        for i in range(len(starts))) if b]


def _parse_block(g: _Grid, sheet: str, top: int, end: int) -> Batch | None:
    hdr = g.label_row(top, end, HEADER_MARK)
    if hdr is None:
        return None

    # --- header row: separate date columns from label columns ---------------
    date_cols, dates = [], []
    for c in range(1, 40):
        v = g.at(hdr, c)
        if _is_date(v):
            date_cols.append(c)
            dates.append(v.date() if isinstance(v, datetime) else v)

    start = _paired(g, top, hdr - 1, "batch start date", raw=True)

    batch = Batch(
        sheet=sheet, top=top,
        batch_code=_paired(g, top, hdr - 1, "batch code"),
        faculty=_paired(g, top, hdr - 1, "faculty"),
        days=_paired(g, top, hdr - 1, "days"),
        time=_paired(g, top, hdr - 1, "time"),
        month=_paired(g, top, hdr - 1, "month"),
        start_date=start.date() if isinstance(start, datetime) else start,
        dates=dates, date_cols=date_cols,
    )

    id_col = g.col_of(hdr, "enrollment #")
    name_col = g.col_of(hdr, "student name")
    course_col = g.col_of(hdr, "course")
    if id_col is None or name_col is None:
        return None

    # The 'S. No' column is pre-numbered past the last student, so it cannot
    # terminate the scan. Stop on a blank enrollment AND name instead.
    for r in range(hdr + 1, end + 1):
        sid, nm = _txt(g.at(r, id_col)), _txt(g.at(r, name_col))
        if not sid and not nm:
            if batch.students:
                break
            continue
        batch.students.append(Student(
            row=r, student_id=sid, name=nm,
            course=_txt(g.at(r, course_col)) if course_col else "",
            marks=[_txt(g.at(r, c)).upper() for c in date_cols],
        ))

    last = batch.students[-1].row if batch.students else hdr
    batch.actual = [_as_int(v) for v in _footer_values(g, last, end, FOOTER_ACTUAL, date_cols)]
    batch.planned = [_as_int(v) for v in _footer_values(g, last, end, FOOTER_CPC, date_cols)]
    batch.modules = [_txt(v) for v in _footer_values(g, last, end, FOOTER_MODULE, date_cols)]
    return batch


def load(path: str | Path) -> list[Batch]:
    wb = load_workbook(Path(path), data_only=True)
    return [b for ws in wb.worksheets for b in parse_sheet(ws)]
